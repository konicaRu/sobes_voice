import io
import struct
import zipfile

import pytest
from docx import Document

from app.documents import MAX_DOCX_UNCOMPRESSED, MAX_TEXT_CHARS, DocumentError, extract_text, safe_filename


def make_docx(build) -> bytes:
    doc = Document()
    build(doc)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def minimal_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return out


def test_markdown_and_txt_are_read_as_utf8():
    assert extract_text("q.md", "# Заголовок\n1. Вопрос?".encode()) == "# Заголовок\n1. Вопрос?"
    assert extract_text("q.txt", "Вопрос?".encode()) == "Вопрос?"


def test_utf8_bom_is_stripped():
    assert extract_text("q.txt", b"\xef\xbb\xbf" + "Вопрос?".encode()) == "Вопрос?"


def test_cp1251_text_is_decoded():
    assert extract_text("q.txt", "Что такое retention?".encode("cp1251")) == "Что такое retention?"


def test_unsupported_extension_is_rejected():
    with pytest.raises(DocumentError, match="Поддерживаются"):
        extract_text("virus.exe", b"MZ")


def test_too_large_text_is_rejected():
    with pytest.raises(DocumentError, match="слишком большой"):
        extract_text("big.txt", ("а" * (MAX_TEXT_CHARS + 1)).encode())


def test_docx_keeps_headings_lists_nesting_and_tables():
    def build(doc):
        doc.add_paragraph("Вопросы", style="Heading 1")
        doc.add_paragraph("Что такое p-value?", style="List Number")
        doc.add_paragraph("вероятность данных", style="List Bullet 2")
        doc.add_paragraph("Обычный абзац")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "левая"
        table.rows[0].cells[1].text = "правая"

    lines = extract_text("q.docx", make_docx(build)).splitlines()
    assert lines[0] == "# Вопросы"
    assert "- Что такое p-value?" in lines
    assert "  - вероятность данных" in lines
    assert "Обычный абзац" in lines
    assert "левая | правая" in lines


def test_docx_with_huge_uncompressed_content_is_rejected():
    data = make_docx(lambda doc: doc.add_paragraph("текст"))
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as bomb:
        for item in source.infolist():
            bomb.writestr(item, source.read(item.filename))
        bomb.writestr("junk.bin", b"\0" * (MAX_DOCX_UNCOMPRESSED + 1))
    with pytest.raises(DocumentError, match="слишком большой"):
        extract_text("bomb.docx", buffer.getvalue())


def test_lying_zip_header_cannot_cause_large_decompression():
    """zipfile читает запись не дальше заявленного размера, а расхождение ловит по CRC (проверено экспериментом)."""
    data = make_docx(lambda doc: doc.add_paragraph("текст"))
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr("junk.bin", b"\0" * (MAX_DOCX_UNCOMPRESSED + 1))
    raw = bytearray(buffer.getvalue())
    with zipfile.ZipFile(io.BytesIO(bytes(raw))) as archive:
        local_offset = archive.getinfo("junk.bin").header_offset
    central = raw.rfind(b"PK\x01\x02", 0, raw.rfind(b"PK\x05\x06"))  # junk.bin записан последним
    struct.pack_into("<I", raw, local_offset + 22, 100)  # размер в локальном заголовке
    struct.pack_into("<I", raw, central + 24, 100)  # размер в центральном каталоге
    assert "текст" in extract_text("liar.docx", bytes(raw))


def test_corrupted_docx_and_pdf_give_readable_errors():
    with pytest.raises(DocumentError, match="прочитать"):
        extract_text("bad.docx", b"not a zip")
    with pytest.raises(DocumentError, match="прочитать"):
        extract_text("bad.pdf", b"not a pdf")


def test_pdf_extraction_stops_as_soon_as_text_limit_is_exceeded(monkeypatch):
    calls = []

    class Page:
        def extract_text(self):
            calls.append(1)
            return "а" * (MAX_TEXT_CHARS + 1)

    class FakeReader:
        def __init__(self, stream):
            self.is_encrypted = False
            self.pages = [Page() for _ in range(50)]

    monkeypatch.setattr("pypdf.PdfReader", FakeReader)
    with pytest.raises(DocumentError, match="слишком большой"):
        extract_text("big.pdf", b"%PDF")
    assert len(calls) == 1


def test_pdf_text_is_extracted():
    assert "What is p-value" in extract_text("q.pdf", minimal_pdf("What is p-value"))


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("../../evil.md", "evil.md"),
        ("C:\\Users\\x\\резюме.docx", "резюме.docx"),
        ('a:b*c?.txt', "a_b_c_.txt"),
        ("CON.txt", "_CON.txt"),
        ("  .скрытый.md ", "скрытый.md"),
    ],
)
def test_safe_filename(raw, expected):
    assert safe_filename(raw) == expected


def test_safe_filename_rejects_bad_extension_and_empty_name():
    with pytest.raises(DocumentError, match="Поддерживаются"):
        safe_filename("script.py")
    with pytest.raises(DocumentError):
        safe_filename("")
